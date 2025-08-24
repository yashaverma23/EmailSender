import streamlit as st
from docx import Document
import smtplib, ssl, os
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication

# -------------------------
# Helpers
# -------------------------
DEFAULT_REASON = "their focus on data-driven decision making and innovation in analytics."

def load_docx_text(path: str) -> str:
    if not os.path.exists(path):
        return ""
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)

def fill_placeholders(template: str, company: str, recruiter: str) -> str:
    filled = template.replace("[Company Name]", company or "[Company Name]")
    filled = filled.replace("[Recruiter’s Name]", recruiter or "[Recruiter’s Name]")
    filled = filled.replace(
        "[specific reason related to company/role—e.g., focus on data-driven decision making, innovation in analytics, or AI-driven business intelligence]",
        DEFAULT_REASON
    )
    return filled

def send_email_smtp(sender_email, sender_password, to_email, subject, body, attachment_bytes=None, attachment_name="Resume.pdf"):
    msg = MIMEMultipart()
    msg["From"] = sender_email
    msg["To"] = to_email
    msg["Subject"] = subject
    msg.attach(MIMEText(body, "plain"))

    if attachment_bytes:
        part = MIMEApplication(attachment_bytes, _subtype="pdf")
        part.add_header("Content-Disposition", "attachment", filename=attachment_name)
        msg.attach(part)

    context = ssl.create_default_context()
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=context) as server:
        server.login(sender_email, sender_password)
        server.send_message(msg)

# -------------------------
# Theme Toggle (Switch)
# -------------------------
def apply_theme(is_dark: bool):
    if is_dark:
        st.markdown(
            """
            <style>
            body, .reportview-container { background-color: #0e1117; color: #e6eef8; }
            .stButton>button { background-color:#0b5fff; color:white; border-radius:6px; }
            input, textarea { background-color:#1c1f26 !important; color:#e6eef8 !important; }
            </style>
            """, unsafe_allow_html=True
        )
    else:
        st.markdown(
            """
            <style>
            body, .reportview-container { background-color: #ffffff; color: #000000; }
            .stButton>button { background-color:#0b5fff; color:white; border-radius:6px; }
            input, textarea { background-color:#f9f9f9 !important; color:#000000 !important; }
            </style>
            """, unsafe_allow_html=True
        )

# Put theme toggle as switch at top right
col1, col2 = st.columns([9, 1])
with col2:
    dark_mode = st.checkbox("🌙", value=False)
apply_theme(dark_mode)

# -------------------------
# UI
# -------------------------
st.title("📧 Cold Email Sender")

col1, col2 = st.columns(2)
with col1:
    to_email = st.text_input("Recipient Email", placeholder="recruiter@company.com")
    company = st.text_input("Company Name", placeholder="Acme Corp")
with col2:
    recruiter = st.text_input("Recruiter’s Name", placeholder="Priya Sharma")
    subject = st.text_input("Subject", f"Application for Data Analyst Role at {company}" if company else "Application for Data Analyst Role")

uploaded = st.file_uploader("Upload Resume (optional, PDF)", type=["pdf"])

# Load template
template_path = "templates/email.docx"
raw_template = load_docx_text(template_path)

placeholders = {
    "[Company Name]": company,
    "[Recruiter’s Name]": recruiter,
}

# Buttons row
colA, colB = st.columns(2)
with colA:
    preview_clicked = st.button("👀 Preview Email")
with colB:
    send_clicked = st.button("📤 Send Email")

# Preview
if preview_clicked:
    if raw_template:
        body_preview = fill_placeholders(raw_template, company, recruiter)
        st.subheader("📄 Email Preview")
        st.text_area("Generated Email", body_preview, height=450)
    else:
        st.error("❌ Could not load email template. Place email.docx in templates/")

# Send
if send_clicked:
    if not to_email:
        st.error("⚠️ Please enter recipient email.")
    elif not raw_template:
        st.error("❌ Email template missing (templates/email.docx).")
    else:
        try:
            sender_email = st.secrets["EMAIL"]
            sender_password = st.secrets["EMAIL_PASSWORD"]

            body = fill_placeholders(raw_template, company, recruiter)

            # Attachment
            attach_bytes, attach_name = None, "Resume.pdf"
            if uploaded is not None:
                attach_bytes = uploaded.read()
                attach_name = uploaded.name
            else:
                default_resume = "templates/Resume.pdf"
                if os.path.exists(default_resume):
                    with open(default_resume, "rb") as f:
                        attach_bytes = f.read()

            send_email_smtp(sender_email, sender_password, to_email, subject, body, attach_bytes, attach_name)
            st.success(f"✅ Email successfully sent to {to_email}")
        except Exception as e:
            st.error(f"❌ Failed to send email: {e}")
